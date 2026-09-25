#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
FdEA: modelo de fórmula abierta (IAE) frente a modelos paramétricos de supervivencia.

Escenarios:
  - 2020 (ajuste): IAE (2020), Tabla 1 del informe IAE (2025).
  - 2025 (validación fuera de muestra): IAE (2025), Tablas 2 y 5 del informe.

Supuestos (IAE, 2025, Tabla 4): tipo real de las aportaciones = crecimiento real pasado del
PIB; tipo real de las prestaciones = crecimiento real futuro del PIB; pensión revalorizada con
el IPC (constante en términos reales); tablas de pensionistas de la Seguridad Social (unisex),
resumidas por su esperanza de vida a los 65 años (21,14 en 2020 y 21,52 en 2025).

Como las tablas de la Seguridad Social no se publican, el modelo de fórmula abierta usa la
forma de la PER2020_Ind_2ndo.orden (unisex) reescalada para reproducir la e65 del IAE.
Los modelos paramétricos se estiman sobre esa tabla en 2020 y, para 2025, solo se actualiza
su parámetro de escala con la nueva e65 (un único dato).

Uso:
    python fdea_weibull_iae.py            # necesita per2020_ind_2orden.csv en la misma carpeta
Salidas: fdea_resultados.xlsx, mallas/*.png, resumen.json
"""
import json
import os
from dataclasses import dataclass, field

import numpy as np
import pandas as pd
from scipy.integrate import quad
from scipy.optimize import brentq, minimize
from scipy.special import gamma as G

ANIOS = list(range(33, 44))

# ------------------------------------------------------------------ datos IAE
EDADES_20 = list(range(60, 71))
CAEJ_20 = np.array([  # Tabla 2 del paper (IAE, 2020)
    [0.565, 0.565, 0.617, 0.687, 0.774, 0.862, 0.932, 0.959, 0.987, 1.014, 1.042],
    [0.597, 0.668, 0.668, 0.704, 0.793, 0.883, 0.954, 0.982, 1.009, 1.037, 1.064],
    [0.611, 0.684, 0.757, 0.794, 0.812, 0.904, 0.977, 1.005, 1.032, 1.060, 1.087],
    [0.625, 0.700, 0.775, 0.850, 0.906, 0.925, 1.000, 1.028, 1.055, 1.083, 1.110],
    [0.625, 0.700, 0.775, 0.850, 0.925, 1.000, 1.000, 1.028, 1.055, 1.083, 1.110],
    [0.625, 0.700, 0.775, 0.850, 0.925, 1.000, 1.028, 1.028, 1.055, 1.083, 1.110],
    [0.650, 0.720, 0.790, 0.860, 0.930, 1.000, 1.040, 1.055, 1.055, 1.083, 1.110],
    [0.650, 0.720, 0.790, 0.860, 0.930, 1.000, 1.040, 1.080, 1.083, 1.110, 1.110],
    [0.650, 0.720, 0.790, 0.860, 0.930, 1.000, 1.040, 1.080, 1.120, 1.110, 1.138],
    [0.675, 0.740, 0.805, 0.870, 0.935, 1.000, 1.040, 1.080, 1.120, 1.160, 1.138],
    [0.675, 0.740, 0.805, 0.870, 0.935, 1.000, 1.040, 1.080, 1.120, 1.160, 1.200]])
FDEA_IAE_20 = np.array([
    [1.43, 1.39, 1.48, 1.59, 1.74, 1.87, 1.95, 1.94, 1.92, 1.90, 1.87],
    [1.44, 1.57, 1.53, 1.56, 1.70, 1.83, 1.91, 1.90, 1.88, 1.85, 1.82],
    [1.41, 1.54, 1.65, 1.68, 1.67, 1.79, 1.87, 1.85, 1.83, 1.81, 1.78],
    [1.38, 1.51, 1.62, 1.72, 1.78, 1.75, 1.83, 1.81, 1.79, 1.77, 1.74],
    [1.33, 1.44, 1.55, 1.65, 1.74, 1.82, 1.75, 1.74, 1.72, 1.69, 1.67],
    [1.27, 1.38, 1.49, 1.58, 1.67, 1.74, 1.73, 1.67, 1.65, 1.63, 1.60],
    [1.27, 1.37, 1.46, 1.54, 1.61, 1.67, 1.68, 1.64, 1.58, 1.56, 1.54],
    [1.22, 1.31, 1.40, 1.48, 1.55, 1.61, 1.62, 1.62, 1.56, 1.54, 1.48],
    [1.17, 1.26, 1.35, 1.42, 1.49, 1.55, 1.55, 1.55, 1.55, 1.48, 1.45],
    [1.17, 1.25, 1.32, 1.38, 1.44, 1.49, 1.49, 1.49, 1.49, 1.48, 1.40],
    [1.12, 1.20, 1.27, 1.33, 1.38, 1.43, 1.43, 1.43, 1.43, 1.43, 1.41]])
EDADES_25 = list(range(63, 71))
TS_25 = np.array([  # IAE (2025), Tabla 2 (%)
    [57.97, 65.33, 78.52, 87.97, 92.02, 96.02, 100.02, 108.02],
    [59.41, 66.95, 80.47, 90.15, 94.30, 98.30, 102.30, 110.30],
    [64.71, 68.57, 82.41, 92.33, 96.58, 100.58, 104.58, 112.58],
    [74.15, 74.15, 84.36, 94.51, 98.86, 102.86, 106.86, 114.86],
    [79.00, 93.23, 93.23, 95.60, 100.00, 104.00, 108.00, 116.00],
    [79.00, 94.50, 96.48, 96.48, 100.00, 104.00, 108.00, 116.00],
    [81.00, 94.75, 100.00, 104.00, 104.00, 104.00, 108.00, 116.00],
    [81.00, 94.75, 100.00, 104.00, 108.00, 108.00, 108.00, 116.00],
    [81.00, 94.75, 100.00, 104.00, 108.00, 112.00, 112.00, 116.00],
    [83.00, 95.00, 100.00, 104.00, 108.00, 112.00, 116.00, 116.00],
    [83.00, 95.00, 100.00, 104.00, 108.00, 112.00, 116.00, 120.00]]) / 100
FDEA_IAE_25 = np.array([  # IAE (2025), Tabla 5
    [1.39, 1.52, 1.76, 1.90, 1.92, 1.97, 2.02, 2.19],
    [1.37, 1.49, 1.73, 1.87, 1.89, 1.93, 1.98, 2.14],
    [1.43, 1.47, 1.70, 1.84, 1.85, 1.90, 1.94, 2.09],
    [1.58, 1.52, 1.67, 1.81, 1.82, 1.86, 1.90, 2.04],
    [1.61, 1.84, 1.78, 1.76, 1.77, 1.81, 1.84, 1.97],
    [1.55, 1.79, 1.77, 1.70, 1.70, 1.73, 1.76, 1.89],
    [1.53, 1.73, 1.76, 1.80, 1.73, 1.66, 1.69, 1.81],
    [1.47, 1.66, 1.69, 1.73, 1.76, 1.69, 1.62, 1.73],
    [1.41, 1.60, 1.63, 1.66, 1.69, 1.71, 1.64, 1.66],
    [1.39, 1.54, 1.57, 1.60, 1.62, 1.65, 1.67, 1.60],
    [1.34, 1.49, 1.51, 1.54, 1.57, 1.59, 1.60, 1.62]])


def tc_2020(anio):
    return 0.1615


def tc_2025(anio):
    return {2023: 0.1672, 2024: 0.1678}.get(anio, 0.1637 if anio <= 2022 else 0.1683)


@dataclass
class Escenario:
    nombre: str
    anio_jub: int
    edades: list
    caej: np.ndarray
    fdea_iae: np.ndarray
    tc: callable
    r_aport: float      # crecimiento real pasado del PIB
    i_prest: float      # crecimiento real futuro del PIB
    e65: float          # esperanza de vida a los 65 (tablas SS, unisex)
    g_sal: float = 0.0022   # crecimiento real de las bases medias (IAE 2025: 0,22 %)
    pagas: int = 14
    exencion: bool = False  # exención de cuotas tras la edad ordinaria (IAE 2025)
    edad_max: int = 120


ESC20 = Escenario("2020", 2020, EDADES_20, CAEJ_20, FDEA_IAE_20, tc_2020, 0.0248, 0.0159, 21.14)
# g_sal 2025: el IAE actualiza la serie histórica de bases medias 2020-2024, que no publica.
# Se fija con el paso (a) de su Tabla 7 (+2,0 % en el FdEA medio), no con la malla.
ESC25 = Escenario("2025", 2025, EDADES_25, TS_25, FDEA_IAE_25, tc_2025, 0.0224, 0.0122, 21.52,
                  g_sal=-0.0013, exencion=True)


def mensual(x):
    return (1 + x) ** (1 / 12) - 1


# ------------------------------------------------------------ aportaciones y BR
def aportaciones_y_br(anios, edad, esc: Escenario):
    """Valores en términos reales; el FdEA no depende del nivel salarial."""
    n = 12 * anios
    t = np.arange(n)
    meses_antes = n - t                                   # meses hasta la jubilación
    base = (1 + mensual(esc.g_sal)) ** (t - (n - 1))      # última base = 1
    anio_cal = esc.anio_jub - np.ceil(meses_antes / 12).astype(int)
    tc = np.array([esc.tc(a) for a in anio_cal])
    if esc.exencion:
        edad_ord = 65 if anios >= 38.25 else 66 + 8 / 12
        edad_mes = edad - meses_antes / 12
        tc = np.where(edad_mes >= edad_ord, 0.0, tc)
    vaa_aport = np.sum(base * tc * (1 + mensual(esc.r_aport)) ** meses_antes)
    br = base[-300:].sum() / 350
    return vaa_aport, br


def fdea_desde_renta(renta_anual_por_edad, esc: Escenario):
    """renta_anual_por_edad[e] = valor actual de 1 u.m. anual vitalicia desde la edad e."""
    out = np.zeros((len(ANIOS), len(esc.edades)))
    for j, e in enumerate(esc.edades):
        for i, a in enumerate(ANIOS):
            aport, br = aportaciones_y_br(a, e, esc)
            out[i, j] = br * esc.caej[i, j] * esc.pagas * renta_anual_por_edad[e] / aport
    return out


def delta(esc):
    return np.log(1 + esc.i_prest)  # pensión constante en términos reales


# ---------------------------------------------- tabla (fórmula abierta, PER2020)
def cargar_per2020(ruta):
    df = pd.read_csv(ruta, sep=";", decimal=",")
    df.columns = [c.strip().lower() for c in df.columns]
    return df.set_index("edad").sort_index()


def kpx_tabla(per, edad, anio_jub, c=1.0, edad_max=120):
    """Supervivencia anual generacional, unisex 50/50, con qx multiplicados por c."""
    def sx(s):
        q = [min(c * per.loc[x, f"qx_base_{s}"] *
                 np.exp(-per.loc[x, f"lambda_{s}"] * (anio_jub + k - 2012)), 1.0)
             for k, x in enumerate(range(edad, edad_max))]
        return np.concatenate([[1.0], np.cumprod(1 - np.array(q))])
    return 0.5 * sx("h") + 0.5 * sx("m")


def mensualizar(kpx):
    k = np.arange(len(kpx))
    tm = np.arange(12 * (len(kpx) - 1) + 1) / 12
    return tm, np.interp(tm, k, kpx)


def e_x(tm, p):
    return np.trapezoid(p, tm) if hasattr(np, "trapezoid") else np.trapz(p, tm)


def calibrar_tabla(per, esc):
    f = lambda c: e_x(*mensualizar(kpx_tabla(per, 65, esc.anio_jub, c))) - esc.e65
    return brentq(f, 0.2, 6)


def rentas_tabla(per, esc, c):
    d = delta(esc)
    out = {}
    for e in esc.edades:
        tm, p = mensualizar(kpx_tabla(per, e, esc.anio_jub, c))
        out[e] = np.sum(np.exp(-d * tm) * p) / 12   # pagos mensuales anticipados
    return out


# ------------------------------------------------------- leyes paramétricas
# Cada ley define tpx(x, t) (t en años) y un parámetro de escala que se reajusta con e65.
class WeibullPaper:
    """Ec. 16-17 del paper: p(t) = exp(-(t / a(edad))^b), t en meses."""
    nombre, n_par = "Weibull (Ec. 16-17)", 4

    def __init__(self, a0=300., a1=0., a2=0., b=3.):
        self.th = np.array([a0, a1, a2, b])
        self.phi = 1.0  # factor de escala común (actualización por e65)

    def a(self, e):
        a0, a1, a2, _ = self.th
        return self.phi * (a0 + a1 * (e - 60) + a2 * (e - 60) ** 2)

    def tpx(self, x, t):
        return np.exp(-(np.asarray(t) * 12 / self.a(x)) ** self.th[3])

    def e65(self):  # fórmula cerrada: E[T] = a(65)·Γ(1+1/b), en meses
        return self.a(65) * G(1 + 1 / self.th[3]) / 12

    def reescalar(self, e65):
        self.phi *= e65 / self.e65()


class WeibullEdad:
    """S(x) = exp(-(x/lam)^k); tpx = S(x+t)/S(x)."""
    nombre, n_par = "Weibull edad alcanzada", 2

    def __init__(self, lam=90., k=10.):
        self.th = np.array([lam, k])

    def tpx(self, x, t):
        lam, k = self.th
        return np.exp(-(((x + t) / lam) ** k - (x / lam) ** k))

    def reescalar(self, e65):
        k = self.th[1]
        self.th[0] = brentq(lambda l: quad(lambda t: np.exp(-(((65 + t) / l) ** k - (65 / l) ** k)),
                                           0, 60)[0] - e65, 60, 130)


class Gompertz:
    """mu(x) = B e^{c x}."""
    nombre, n_par = "Gompertz", 2

    def __init__(self, lB=-10., c=0.1):
        self.th = np.array([lB, c])

    def tpx(self, x, t):
        lB, c = self.th
        return np.exp(-np.exp(lB) / c * np.exp(c * x) * (np.exp(c * np.asarray(t)) - 1))

    def reescalar(self, e65):
        c = self.th[1]
        self.th[0] = brentq(lambda lB: quad(lambda t: np.exp(-np.exp(lB) / c * np.exp(c * 65) *
                                                              (np.exp(c * t) - 1)), 0, 60)[0] - e65,
                            -30, 0)


class Makeham(Gompertz):
    """mu(x) = A + B e^{c x}."""
    nombre, n_par = "Gompertz-Makeham", 3

    def __init__(self, A=1e-3, lB=-10., c=0.1):
        self.th = np.array([A, lB, c])

    def tpx(self, x, t):
        A, lB, c = self.th
        t = np.asarray(t)
        return np.exp(-A * t - np.exp(lB) / c * np.exp(c * x) * (np.exp(c * t) - 1))

    def reescalar(self, e65):
        A, _, c = self.th
        f = lambda lB: quad(lambda t: np.exp(-A * t - np.exp(lB) / c * np.exp(c * 65) *
                                             (np.exp(c * t) - 1)), 0, 60)[0] - e65
        self.th[1] = brentq(f, -30, 0)


class Exponencial:
    """mu constante = 1/e65 (referencia ingenua, 1 parámetro)."""
    nombre, n_par = "Exponencial", 1

    def __init__(self, mu=0.05):
        self.th = np.array([mu])

    def tpx(self, x, t):
        return np.exp(-self.th[0] * np.asarray(t))

    def reescalar(self, e65):
        self.th[0] = 1 / e65


def estimar(ley, objetivo, x0, bounds):
    """Mínimos cuadrados entre tpx de la ley y la supervivencia mensual de la tabla (60-70)."""
    def sce(th):
        ley.th = np.array(th, dtype=float)
        return sum(np.sum((ley.tpx(e, tm) - p) ** 2) for e, (tm, p) in objetivo.items())
    res = minimize(sce, x0, method="L-BFGS-B", bounds=bounds)
    ley.th = res.x
    return ley


def rentas_ley(ley, esc):
    d = delta(esc)
    T = esc.edad_max
    return {e: quad(lambda t: np.exp(-d * t) * ley.tpx(e, t), 0, T - e, limit=400)[0]
            for e in esc.edades}


# ------------------------------------------------------------- métricas y mallas
def metricas(ref, mod):
    d = mod - ref
    return dict(MAPE=100 * np.mean(np.abs(d / ref)), MAE=np.mean(np.abs(d)),
                max_abs=np.max(np.abs(d)), max_rel=100 * np.max(np.abs(d / ref)),
                sesgo=np.mean(d), pct_le_005=100 * np.mean(np.abs(d) <= 0.05 + 1e-9),
                corr=np.corrcoef(ref.ravel(), mod.ravel())[0, 1])


ROJO, AMARILLO, VERDE = "F8696B", "FFEB84", "63BE7B"


def _mezcla(c1, c2, f):
    a = [int(c1[i:i + 2], 16) for i in (0, 2, 4)]
    b = [int(c2[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x + (y - x) * f):02X}" for x, y in zip(a, b))


def colores_escala(m, invertir=False):
    """Escala de 3 colores de Excel (rojo-amarillo-verde, mediana como punto medio), como el IAE."""
    lo, hi = (VERDE, ROJO) if invertir else (ROJO, VERDE)
    vmin, vmed, vmax = np.min(m), np.median(m), np.max(m)
    out = np.empty(m.shape, dtype=object)
    for idx, v in np.ndenumerate(m):
        if v <= vmed:
            out[idx] = _mezcla(lo, AMARILLO, 0 if vmed == vmin else (v - vmin) / (vmed - vmin))
        else:
            out[idx] = _mezcla(AMARILLO, hi, 0 if vmax == vmed else (v - vmed) / (vmax - vmed))
    return out


def malla_png(m, edades, titulo, ruta, invertir=False, colorear=None):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    cols = colores_escala(m if colorear is None else colorear, invertir)
    fig, ax = plt.subplots(figsize=(0.75 * len(edades) + 1.5, 4.8))
    for i in range(m.shape[0]):
        for j in range(m.shape[1]):
            ax.add_patch(plt.Rectangle((j, i), 1, 1, facecolor="#" + cols[i, j], edgecolor="black", lw=0.5))
            v = 0.0 if abs(m[i, j]) < 0.005 else m[i, j]
            ax.text(j + 0.5, i + 0.5, f"{v:.2f}".replace(".", ",").replace("-", "−"),
                    ha="center", va="center", fontsize=8)
    ax.set_xlim(0, m.shape[1]); ax.set_ylim(m.shape[0], 0)
    ax.set_xticks(np.arange(m.shape[1]) + 0.5); ax.set_xticklabels(edades)
    ax.set_yticks(np.arange(m.shape[0]) + 0.5); ax.set_yticklabels(ANIOS)
    ax.xaxis.tick_top(); ax.xaxis.set_label_position("top")
    ax.set_xlabel("Edad de jubilación"); ax.set_ylabel("Años cotizados"); ax.tick_params(length=0)
    ax.set_title(titulo, fontsize=10, pad=28)
    fig.tight_layout(); fig.savefig(ruta, dpi=200); plt.close(fig)


def hoja_malla(ws, m, edades, invertir=False, colorear=None):
    from openpyxl.styles import PatternFill, Alignment, Border, Side, Font
    s = Side(style="thin", color="000000"); b = Border(left=s, right=s, top=s, bottom=s)
    cols = colores_escala(m if colorear is None else colorear, invertir)
    ws.cell(1, 1, "Años cotizados").font = Font(bold=True)
    ws.cell(1, 2, "Edad de jubilación").font = Font(bold=True)
    for j, e in enumerate(edades):
        c = ws.cell(2, j + 2, e); c.border = b; c.alignment = Alignment(horizontal="center")
    for i, a in enumerate(ANIOS):
        c = ws.cell(3 + i, 1, a); c.border = b; c.alignment = Alignment(horizontal="center")
        for j in range(len(edades)):
            c = ws.cell(3 + i, j + 2, round(float(m[i, j]), 2))
            c.number_format = "0.00"; c.border = b; c.alignment = Alignment(horizontal="center")
            c.fill = PatternFill("solid", start_color=cols[i, j], end_color=cols[i, j])
    ws.column_dimensions["A"].width = 15


# ------------------------------------------------------------------------ main
def main(ruta_per="per2020_ind_2orden.csv", salida="fdea_resultados.xlsx"):
    per = cargar_per2020(ruta_per)
    res, tablas, mets = {}, {}, []

    # 1) Tabla PER2020 reescalada a la e65 del IAE (fórmula abierta)
    c20, c25 = calibrar_tabla(per, ESC20), calibrar_tabla(per, ESC25)
    res["factor_qx"] = {"2020": c20, "2025": c25}
    for esc, c in [(ESC20, c20), (ESC25, c25)]:
        tablas[f"Abierta_{esc.nombre}"] = fdea_desde_renta(rentas_tabla(per, esc, c), esc)

    # 2) Leyes paramétricas: estimadas sobre la tabla 2020; en 2025 solo cambia la escala (e65)
    objetivo = {e: mensualizar(kpx_tabla(per, e, 2020, c20)) for e in EDADES_20}
    leyes = [
        (WeibullPaper(), [400., -10., 0., 3.], [(1, None), (None, None), (None, None), (0.05, 30)]),
        (WeibullEdad(), [90., 10.], [(60, 130), (1.01, 40)]),
        (Gompertz(), [-10., 0.1], [(-30, 0), (0.001, 0.5)]),
        (Makeham(), [1e-3, -10., 0.1], [(0, 0.05), (-30, 0), (0.001, 0.5)]),
        (Exponencial(), [0.05], [(1e-4, 1)]),
    ]
    res["parametros"] = {}
    for ley, x0, bnd in leyes:
        estimar(ley, objetivo, x0, bnd)
        ley.reescalar(ESC20.e65)   # todas las leyes reproducen exactamente la e65 del IAE
        th20 = ley.th.copy(); phi20 = getattr(ley, "phi", None)
        tablas[f"{ley.nombre}_2020"] = fdea_desde_renta(rentas_ley(ley, ESC20), ESC20)
        ley.reescalar(ESC25.e65)
        tablas[f"{ley.nombre}_2025"] = fdea_desde_renta(rentas_ley(ley, ESC25), ESC25)
        res["parametros"][ley.nombre] = {"2020": th20.tolist(),
                                         "2025": ley.th.tolist(),
                                         "phi_2025": getattr(ley, "phi", None)}

    # 3) Aplicación: sensibilidad del FdEA 2025 a la longevidad (Weibull Ec. 16-17)
    wp = leyes[0][0]
    sens = {}
    for e65 in [21.52, 22.52, 22.63, 23.11]:
        wp.reescalar(e65)
        sens[e65] = fdea_desde_renta(rentas_ley(wp, ESC25), ESC25)
    wp.reescalar(ESC25.e65)
    tablas["Sens_+1anio_2025"] = sens[22.52] - sens[21.52]
    res["sensibilidad"] = {str(k): float(np.mean(v)) for k, v in sens.items()}
    res["elasticidad_e65"] = float(np.mean((sens[22.52] / sens[21.52] - 1) / (1 / 21.52)))

    # 4) Métricas frente al IAE
    for nombre, m in list(tablas.items()):
        if nombre.startswith("Sens"):
            continue
        esc = ESC20 if nombre.endswith("2020") else ESC25
        d = metricas(esc.fdea_iae, m); d.update(modelo=nombre.rsplit("_", 1)[0], escenario=esc.nombre)
        mets.append(d)
    dfm = pd.DataFrame(mets)[["modelo", "escenario", "MAPE", "MAE", "max_abs", "max_rel",
                              "sesgo", "pct_le_005", "corr"]].round(4)
    res["metricas"] = dfm.to_dict(orient="records")
    print(dfm.to_string(index=False))

    # 5) Exportación
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "metricas"
    ws.append(list(dfm.columns))
    for r in dfm.itertuples(index=False):
        ws.append(list(r))
    os.makedirs("mallas", exist_ok=True)
    todas = {"IAE_2020": FDEA_IAE_20, "IAE_2025": FDEA_IAE_25, **tablas}
    for nombre, m in todas.items():
        edades = EDADES_20 if nombre.endswith("2020") else EDADES_25
        seguro = nombre.replace(" ", "_").replace("(", "").replace(")", "").replace(".", "")
        hoja_malla(wb.create_sheet(seguro[:31]), m, edades)
        malla_png(m, edades, nombre.replace("_", " "), f"mallas/{seguro}.png")
        if not nombre.startswith(("IAE", "Sens")):
            ref = FDEA_IAE_20 if nombre.endswith("2020") else FDEA_IAE_25
            d = ref - m
            hoja_malla(wb.create_sheet(("dif_" + seguro)[:31]), d, edades, invertir=True, colorear=np.abs(d))
            malla_png(d, edades, "Diferencia IAE − " + nombre.replace("_", " "),
                      f"mallas/dif_{seguro}.png", invertir=True, colorear=np.abs(d))
    wb.save(salida)
    res["tablas"] = {k: np.round(v, 4).tolist() for k, v in todas.items()}
    with open("resumen.json", "w") as fh:
        json.dump(res, fh, indent=1, ensure_ascii=False)
    print(json.dumps({k: res[k] for k in ("factor_qx", "parametros", "sensibilidad",
                                          "elasticidad_e65")}, indent=1))


if __name__ == "__main__":
    main()
